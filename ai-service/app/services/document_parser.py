"""Multi-format document parsing service for Custom Chatbots.

Supports:
- PDF (.pdf) via pypdf with page preservation
- Word (.docx) via python-docx with headings, tables, and paragraphs
- PowerPoint (.pptx) via python-pptx with slides and speaker notes
- Plain text & Markdown (.txt, .md)
- Tabular CSV (.csv) converted to structured Markdown tables
- Images (.png, .jpg, .jpeg, .webp) via Gemini Vision multimodal OCR
"""

import base64
import csv
import io
import logging
import os
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional

import httpx
from PIL import Image

logger = logging.getLogger(__name__)


@dataclass
class DocumentSection:
    """A distinct structural subsection of a document (e.g. Page 1, Slide 3, Table, Heading)."""

    title: str
    content: str
    section_number: int
    metadata: Dict[str, Any] = field(default_factory=dict)


@dataclass
class ParsedDocumentResult:
    """Standardized output produced from parsing any supported document."""

    filename: str
    file_type: str
    title: Optional[str]
    total_pages_or_slides: int
    sections: List[DocumentSection]
    full_text: str
    char_count: int
    metadata: Dict[str, Any] = field(default_factory=dict)


class DocumentParsingError(Exception):
    """Raised when parsing or extracting text from a document fails."""
    pass


class DocumentParserService:
    """Unified service for parsing and extracting structured text from multiple document formats."""

    GEMINI_API_URL = "https://generativelanguage.googleapis.com/v1beta/models"
    DEFAULT_VISION_MODEL = "gemini-flash-latest"

    # Supported MIME types and file extensions
    SUPPORTED_EXTENSIONS = {
        ".pdf": "application/pdf",
        ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        ".txt": "text/plain",
        ".md": "text/markdown",
        ".csv": "text/csv",
        ".png": "image/png",
        ".jpg": "image/jpeg",
        ".jpeg": "image/jpeg",
        ".webp": "image/webp",
    }

    def is_supported(self, filename: str, mime_type: Optional[str] = None) -> bool:
        """Checks if a file extension or MIME type is supported."""
        ext = Path(filename).suffix.lower()
        if ext in self.SUPPORTED_EXTENSIONS:
            return True
        if mime_type and mime_type.lower() in self.SUPPORTED_EXTENSIONS.values():
            return True
        return False

    def detect_file_type(self, filename: str, mime_type: Optional[str] = None) -> str:
        """Determines the normalized file type (pdf, docx, pptx, txt, md, csv, image)."""
        ext = Path(filename).suffix.lower()
        if ext == ".pdf" or mime_type == "application/pdf":
            return "pdf"
        if ext in (".docx", ".doc") or (mime_type and "wordprocessingml" in mime_type):
            return "docx"
        if ext in (".pptx", ".ppt") or (mime_type and "presentationml" in mime_type):
            return "pptx"
        if ext == ".csv" or mime_type == "text/csv":
            return "csv"
        if ext == ".md" or mime_type == "text/markdown":
            return "md"
        if ext in (".txt", ".text") or mime_type == "text/plain":
            return "txt"
        if ext in (".png", ".jpg", ".jpeg", ".webp") or (mime_type and mime_type.startswith("image/")):
            return "image"
        return "txt"

    async def parse_document(
        self,
        file_bytes: bytes,
        filename: str,
        mime_type: Optional[str] = None,
        api_key: Optional[str] = None,
        model_name: Optional[str] = None,
    ) -> ParsedDocumentResult:
        """Parses any supported document format into a structured ParsedDocumentResult."""
        if not file_bytes:
            raise DocumentParsingError(f"Uploaded file '{filename}' is empty (0 bytes).")

        file_type = self.detect_file_type(filename, mime_type)
        logger.info(f"Parsing document '{filename}' as file_type '{file_type}' ({len(file_bytes)} bytes)")

        try:
            if file_type == "pdf":
                return self.parse_pdf(file_bytes, filename)
            elif file_type == "docx":
                return self.parse_docx(file_bytes, filename)
            elif file_type == "pptx":
                return self.parse_pptx(file_bytes, filename)
            elif file_type == "csv":
                return self.parse_csv(file_bytes, filename)
            elif file_type in ("txt", "md"):
                return self.parse_text(file_bytes, filename, file_type)
            elif file_type == "image":
                return await self.parse_image(file_bytes, filename, mime_type, api_key, model_name)
            else:
                return self.parse_text(file_bytes, filename, "txt")
        except DocumentParsingError:
            raise
        except Exception as exc:
            logger.error(f"Error parsing document '{filename}': {exc}", exc_info=True)
            raise DocumentParsingError(f"Failed to parse document '{filename}': {str(exc)}") from exc

    # =========================================================================
    # PDF PARSER (pypdf)
    # =========================================================================
    def parse_pdf(self, file_bytes: bytes, filename: str) -> ParsedDocumentResult:
        """Extracts text page-by-page from a PDF document preserving page numbers."""
        import pypdf

        stream = io.BytesIO(file_bytes)
        try:
            reader = pypdf.PdfReader(stream)
        except Exception as exc:
            raise DocumentParsingError(f"Invalid or corrupted PDF file: {exc}")

        total_pages = len(reader.pages)
        if total_pages == 0:
            raise DocumentParsingError("PDF contains 0 pages.")

        sections: List[DocumentSection] = []
        full_text_parts: List[str] = []

        doc_title: Optional[str] = None
        if reader.metadata and reader.metadata.title:
            doc_title = str(reader.metadata.title).strip() or None

        for page_num, page in enumerate(reader.pages, start=1):
            try:
                page_text = page.extract_text() or ""
            except Exception as page_err:
                logger.warning(f"Error extracting text from page {page_num} in '{filename}': {page_err}")
                page_text = ""

            cleaned_text = page_text.strip()
            if cleaned_text:
                section_title = f"Page {page_num}"
                sections.append(
                    DocumentSection(
                        title=section_title,
                        content=cleaned_text,
                        section_number=page_num,
                        metadata={"page": page_num, "total_pages": total_pages},
                    )
                )
                full_text_parts.append(f"## {section_title}\n\n{cleaned_text}")

        full_text = "\n\n".join(full_text_parts).strip()
        if not full_text:
            raise DocumentParsingError(
                f"PDF '{filename}' contained no selectable text. It may contain scanned image pages without OCR."
            )

        return ParsedDocumentResult(
            filename=filename,
            file_type="pdf",
            title=doc_title or Path(filename).stem,
            total_pages_or_slides=total_pages,
            sections=sections,
            full_text=full_text,
            char_count=len(full_text),
            metadata={"total_pages": total_pages, "author": getattr(reader.metadata, "author", None)},
        )

    # =========================================================================
    # WORD PARSER (.docx via python-docx)
    # =========================================================================
    def parse_docx(self, file_bytes: bytes, filename: str) -> ParsedDocumentResult:
        """Extracts text, headings, bullet points, and tables from Word docx."""
        import docx

        stream = io.BytesIO(file_bytes)
        try:
            doc = docx.Document(stream)
        except Exception as exc:
            raise DocumentParsingError(f"Invalid or corrupted Word document: {exc}")

        sections: List[DocumentSection] = []
        full_text_parts: List[str] = []

        current_heading = "Overview"
        current_section_lines: List[str] = []
        section_counter = 1

        doc_title: Optional[str] = None

        def flush_section():
            nonlocal section_counter, current_section_lines
            content = "\n\n".join(current_section_lines).strip()
            if content:
                sections.append(
                    DocumentSection(
                        title=current_heading,
                        content=content,
                        section_number=section_counter,
                        metadata={"section_name": current_heading},
                    )
                )
                full_text_parts.append(f"## {current_heading}\n\n{content}")
                section_counter += 1
            current_section_lines = []

        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue

            style_name = getattr(p.style, "name", "")
            if style_name.startswith("Heading") or style_name in ("Title", "Subtitle"):
                flush_section()
                current_heading = text
                if not doc_title and style_name in ("Heading 1", "Title"):
                    doc_title = text
            else:
                current_section_lines.append(text)

        # Extract structured Markdown tables
        for table_idx, table in enumerate(doc.tables, start=1):
            table_rows: List[List[str]] = []
            for row in table.rows:
                cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                table_rows.append(cells)

            if table_rows:
                md_table_lines: List[str] = []
                headers = table_rows[0]
                md_table_lines.append("| " + " | ".join(headers) + " |")
                md_table_lines.append("| " + " | ".join(["---"] * len(headers)) + " |")
                for data_row in table_rows[1:]:
                    # Pad or trim to match header length
                    padded = data_row[: len(headers)] + [""] * max(0, len(headers) - len(data_row))
                    md_table_lines.append("| " + " | ".join(padded) + " |")
                table_str = "\n".join(md_table_lines)
                current_section_lines.append(f"### Table {table_idx}\n{table_str}")

        flush_section()

        full_text = "\n\n".join(full_text_parts).strip()
        if not full_text:
            raise DocumentParsingError(f"Word document '{filename}' contained no readable text or tables.")

        return ParsedDocumentResult(
            filename=filename,
            file_type="docx",
            title=doc_title or Path(filename).stem,
            total_pages_or_slides=len(sections) or 1,
            sections=sections,
            full_text=full_text,
            char_count=len(full_text),
            metadata={"sections_count": len(sections), "tables_count": len(doc.tables)},
        )

    # =========================================================================
    # POWERPOINT PARSER (.pptx via python-pptx)
    # =========================================================================
    def parse_pptx(self, file_bytes: bytes, filename: str) -> ParsedDocumentResult:
        """Extracts text, slide titles, bullet points, and presenter notes from PPTX."""
        import pptx

        stream = io.BytesIO(file_bytes)
        try:
            prs = pptx.Presentation(stream)
        except Exception as exc:
            raise DocumentParsingError(f"Invalid or corrupted PowerPoint presentation: {exc}")

        total_slides = len(prs.slides)
        if total_slides == 0:
            raise DocumentParsingError("Presentation contains 0 slides.")

        sections: List[DocumentSection] = []
        full_text_parts: List[str] = []

        presentation_title: Optional[str] = None

        for slide_idx, slide in enumerate(prs.slides, start=1):
            slide_title = f"Slide {slide_idx}"
            slide_text_lines: List[str] = []

            # Extract shapes and titles
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        txt = paragraph.text.strip()
                        if txt:
                            if shape == slide.shapes.title:
                                slide_title = f"Slide {slide_idx}: {txt}"
                                if not presentation_title and slide_idx == 1:
                                    presentation_title = txt
                            else:
                                slide_text_lines.append(txt)
                elif shape.has_table:
                    table = shape.table
                    for row in table.rows:
                        row_cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                        if any(row_cells):
                            slide_text_lines.append(" | ".join(row_cells))

            # Extract speaker notes
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes_text = slide.notes_slide.notes_text_frame.text.strip()
                if notes_text:
                    slide_text_lines.append(f"Speaker Notes: {notes_text}")

            content = "\n".join(slide_text_lines).strip()
            if content:
                sections.append(
                    DocumentSection(
                        title=slide_title,
                        content=content,
                        section_number=slide_idx,
                        metadata={"slide": slide_idx, "total_slides": total_slides},
                    )
                )
                full_text_parts.append(f"## {slide_title}\n\n{content}")

        full_text = "\n\n".join(full_text_parts).strip()
        if not full_text:
            raise DocumentParsingError(f"PowerPoint '{filename}' contained no text across its {total_slides} slides.")

        return ParsedDocumentResult(
            filename=filename,
            file_type="pptx",
            title=presentation_title or Path(filename).stem,
            total_pages_or_slides=total_slides,
            sections=sections,
            full_text=full_text,
            char_count=len(full_text),
            metadata={"total_slides": total_slides},
        )

    # =========================================================================
    # CSV / SPREADSHEET PARSER
    # =========================================================================
    def parse_csv(self, file_bytes: bytes, filename: str) -> ParsedDocumentResult:
        """Parses CSV into a clean, searchable Markdown table format."""
        text_content = self._decode_bytes(file_bytes)
        reader = csv.reader(io.StringIO(text_content))

        rows: List[List[str]] = [row for row in reader if any(cell.strip() for cell in row)]
        if not rows:
            raise DocumentParsingError(f"CSV file '{filename}' contains no data rows.")

        headers = [c.strip() for c in rows[0]]
        md_lines = ["| " + " | ".join(headers) + " |", "| " + " | ".join(["---"] * len(headers)) + " |"]

        for row in rows[1:]:
            padded = [c.strip().replace("\n", " ") for c in row[: len(headers)]]
            padded += [""] * max(0, len(headers) - len(padded))
            md_lines.append("| " + " | ".join(padded) + " |")

        formatted_table = "\n".join(md_lines)
        section = DocumentSection(
            title=f"Dataset: {Path(filename).stem}",
            content=formatted_table,
            section_number=1,
            metadata={"rows_count": len(rows), "columns_count": len(headers)},
        )

        return ParsedDocumentResult(
            filename=filename,
            file_type="csv",
            title=Path(filename).stem,
            total_pages_or_slides=1,
            sections=[section],
            full_text=formatted_table,
            char_count=len(formatted_table),
            metadata={"rows_count": len(rows), "columns_count": len(headers)},
        )

    # =========================================================================
    # PLAIN TEXT & MARKDOWN PARSER
    # =========================================================================
    def parse_text(self, file_bytes: bytes, filename: str, file_type: str = "txt") -> ParsedDocumentResult:
        """Decodes plain text or Markdown files."""
        text_content = self._decode_bytes(file_bytes).strip()
        if not text_content:
            raise DocumentParsingError(f"Text file '{filename}' is empty.")

        sections: List[DocumentSection] = []
        # Chunk text into logical paragraphs or Markdown headers
        raw_blocks = [b.strip() for b in text_content.split("\n\n") if b.strip()]
        for idx, block in enumerate(raw_blocks, start=1):
            sections.append(
                DocumentSection(
                    title=f"Section {idx}",
                    content=block,
                    section_number=idx,
                    metadata={"block_idx": idx},
                )
            )

        return ParsedDocumentResult(
            filename=filename,
            file_type=file_type,
            title=Path(filename).stem,
            total_pages_or_slides=len(sections) or 1,
            sections=sections,
            full_text=text_content,
            char_count=len(text_content),
            metadata={"blocks_count": len(sections)},
        )

    # =========================================================================
    # IMAGE OCR & VISION PARSER (Gemini Vision)
    # =========================================================================
    async def parse_image(
        self,
        file_bytes: bytes,
        filename: str,
        mime_type: Optional[str] = None,
        api_key: Optional[str] = None,
        model_name: Optional[str] = None,
    ) -> ParsedDocumentResult:
        """Analyzes an image using Gemini Vision to transcribe text and describe visual diagrams."""
        try:
            pil_image = Image.open(io.BytesIO(file_bytes))
            width, height = pil_image.size
            img_format = pil_image.format or "JPEG"
        except Exception as img_err:
            raise DocumentParsingError(f"Invalid or corrupted image file '{filename}': {img_err}")

        # Resolve correct MIME type
        norm_mime = mime_type
        if not norm_mime:
            ext = Path(filename).suffix.lower()
            norm_mime = self.SUPPORTED_EXTENSIONS.get(ext, f"image/{img_format.lower()}")

        effective_key = api_key or os.getenv("GEMINI_API_KEY")
        if not effective_key:
            # Fallback when no Gemini API key is configured
            fallback_text = (
                f"Image Reference: {filename}\n"
                f"Dimensions: {width}x{height} px, Format: {img_format}\n"
                f"(Vision OCR unavailable: please configure GEMINI_API_KEY for automatic image transcription.)"
            )
            section = DocumentSection(
                title=f"Image: {Path(filename).stem}",
                content=fallback_text,
                section_number=1,
                metadata={"width": width, "height": height, "format": img_format},
            )
            return ParsedDocumentResult(
                filename=filename,
                file_type="image",
                title=Path(filename).stem,
                total_pages_or_slides=1,
                sections=[section],
                full_text=fallback_text,
                char_count=len(fallback_text),
                metadata={"width": width, "height": height, "ocr_status": "skipped_no_key"},
            )

        # Call Gemini Vision to extract text and describe diagrams
        b64_data = base64.b64encode(file_bytes).decode("utf-8")
        vision_model = model_name or self.DEFAULT_VISION_MODEL
        target_url = f"{self.GEMINI_API_URL}/{vision_model}:generateContent?key={effective_key}"

        prompt = (
            "You are a factual knowledge extraction engine for an AI assistant. "
            "Examine this image and extract all useful reference information:\n"
            "1. Transcribe all text, numbers, formulas, titles, labels, or captions verbatim.\n"
            "2. If this image contains a table, transcribe it into a clean Markdown table.\n"
            "3. If this image is a diagram, workflow, infographic, or flowchart, explain each step and relation factually.\n"
            "Provide a clean, comprehensive markdown transcription suitable for semantic search and RAG retrieval."
        )

        payload = {
            "contents": [
                {
                    "parts": [
                        {"text": prompt},
                        {
                            "inline_data": {
                                "mime_type": norm_mime,
                                "data": b64_data,
                            }
                        },
                    ]
                }
            ],
            "generationConfig": {
                "temperature": 0.1,
                "maxOutputTokens": 2048,
            },
        }

        try:
            async with httpx.AsyncClient(timeout=45.0) as client:
                res = await client.post(target_url, json=payload)
                if res.status_code == 200:
                    data = res.json()
                    candidates = data.get("candidates", [])
                    extracted_text = ""
                    if candidates:
                        parts = candidates[0].get("content", {}).get("parts", [])
                        extracted_text = "\n".join(p.get("text", "") for p in parts if "text" in p).strip()

                    if not extracted_text:
                        extracted_text = f"Image '{filename}' processed, but no text or visual features were found."

                    section = DocumentSection(
                        title=f"Image Vision OCR: {Path(filename).stem}",
                        content=extracted_text,
                        section_number=1,
                        metadata={"width": width, "height": height, "format": img_format, "ocr": "gemini_vision"},
                    )
                    return ParsedDocumentResult(
                        filename=filename,
                        file_type="image",
                        title=Path(filename).stem,
                        total_pages_or_slides=1,
                        sections=[section],
                        full_text=extracted_text,
                        char_count=len(extracted_text),
                        metadata={"width": width, "height": height, "format": img_format, "ocr_status": "success"},
                    )
                else:
                    logger.warning(f"Gemini Vision call failed with status {res.status_code}: {res.text}")
        except Exception as ocr_err:
            logger.error(f"Image vision extraction error for '{filename}': {ocr_err}")

        # Fallback to metadata if Vision API timed out or errored
        fallback_text = (
            f"Image File: {filename} ({width}x{height} {img_format}). Vision extraction temporarily unavailable."
        )
        return ParsedDocumentResult(
            filename=filename,
            file_type="image",
            title=Path(filename).stem,
            total_pages_or_slides=1,
            sections=[
                DocumentSection(
                    title=f"Image: {Path(filename).stem}",
                    content=fallback_text,
                    section_number=1,
                    metadata={"width": width, "height": height},
                )
            ],
            full_text=fallback_text,
            char_count=len(fallback_text),
            metadata={"width": width, "height": height, "ocr_status": "error_fallback"},
        )

    @staticmethod
    def _decode_bytes(data: bytes) -> str:
        """Safely decodes raw byte content with multiple fallback encodings."""
        for encoding in ("utf-8", "utf-8-sig", "latin-1", "cp1252", "iso-8859-1"):
            try:
                return data.decode(encoding)
            except UnicodeDecodeError:
                continue
        return data.decode("utf-8", errors="replace")


# Singleton instance
document_parser_service = DocumentParserService()
